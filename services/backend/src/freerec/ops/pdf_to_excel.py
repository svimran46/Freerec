"""Extracts tables from a PDF into an Excel workbook, one sheet per page with tables."""

import re
from pathlib import Path

import pdfplumber
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from pydantic import BaseModel

from freerec.ops import InputPath, NoOptions, Op, OpError
from freerec.ops._pdf import check_pdf
from freerec.ops._tools import output_name

_NUMBER = re.compile(r"^-?(\d{1,3}(,\d{3})+|\d+)(\.\d+)?$")


def cell_value(text: str | None) -> str | int | float | None:
    """Turns numbers written in a table ("1,234.50") into real numbers; keeps other text."""
    if text is None:
        return None
    text = text.strip()
    if _NUMBER.match(text):
        plain = text.replace(",", "")
        return float(plain) if "." in plain else int(plain)
    return text or None


def run(inputs: list[InputPath], _options: BaseModel, out_dir: Path) -> list[Path]:
    [item] = inputs
    check_pdf(item)
    workbook = Workbook()
    workbook.remove(workbook.active)
    try:
        with pdfplumber.open(item.path) as pdf:
            for number, page in enumerate(pdf.pages, start=1):
                tables = page.extract_tables()
                if not tables:
                    continue
                sheet = workbook.create_sheet(f"Page {number}")
                row = 1
                for table in tables:
                    for cells in table:
                        for column, text in enumerate(cells, start=1):
                            sheet.cell(row=row, column=column, value=cell_value(text))
                        row += 1
                    row += 1  # A blank row between tables.
                for column in range(1, sheet.max_column + 1):
                    sheet.column_dimensions[get_column_letter(column)].width = 18
            if not workbook.sheetnames:
                # No ruled tables: give the text, one line per row, rather than nothing.
                sheet = workbook.create_sheet("Text")
                row = 1
                for page in pdf.pages:
                    for line in (page.extract_text() or "").splitlines():
                        sheet.cell(row=row, column=1, value=line)
                        row += 1
                if row == 1:
                    raise OpError(f"{item.name} has no text to extract. It may be a scan.")
    except OpError:
        raise
    except Exception as error:
        raise OpError(f"{item.name} couldn't be read. It may be damaged.") from error

    target = out_dir / output_name(item.name, "xlsx")
    workbook.save(target)
    return [target]


OP = Op(slug="pdf-to-excel", accepts=frozenset({"pdf"}), run=run, options=NoOptions)
