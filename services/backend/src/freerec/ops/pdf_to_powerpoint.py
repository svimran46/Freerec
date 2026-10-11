"""Turns each PDF page into a slide holding an image of the page."""

import io
from pathlib import Path

from pydantic import BaseModel

from freerec.ops import InputPath, NoOptions, Op
from freerec.ops._pdf import open_pdf
from freerec.ops._tools import output_name

DPI = 150
EMU_PER_POINT = 12700
# PowerPoint accepts slides from 1 to 56 inches on each side.
MIN_SIDE, MAX_SIDE = 72 * EMU_PER_POINT, 56 * 72 * EMU_PER_POINT


def slide_size(width_pt: float, height_pt: float) -> tuple[int, int]:
    """The slide size in EMU for a page, scaled to fit PowerPoint's limits."""
    width, height = width_pt * EMU_PER_POINT, height_pt * EMU_PER_POINT
    scale = max(min(1.0, MAX_SIDE / width, MAX_SIDE / height), MIN_SIDE / width, MIN_SIDE / height)
    # Pages more than 56 times wider than tall can't meet both limits; the maximum wins.
    scale = min(scale, MAX_SIDE / width, MAX_SIDE / height)
    return round(width * scale), round(height * scale)


def run(inputs: list[InputPath], _options: BaseModel, out_dir: Path) -> list[Path]:
    from pptx import Presentation
    from pptx.util import Emu

    [item] = inputs
    doc = open_pdf(item)
    try:
        deck = Presentation()
        width, height = slide_size(*doc[0].get_size())
        deck.slide_width, deck.slide_height = Emu(width), Emu(height)
        blank = deck.slide_layouts[6]
        for page in doc:
            image = page.render(scale=DPI / 72).to_pil().convert("RGB")
            buffer = io.BytesIO()
            image.save(buffer, format="JPEG", quality=85)
            buffer.seek(0)
            # Pages of a different shape are fitted and centered on the slide.
            page_w, page_h = page.get_size()
            scale = min(width / page_w, height / page_h)
            pic_w, pic_h = round(page_w * scale), round(page_h * scale)
            slide = deck.slides.add_slide(blank)
            slide.shapes.add_picture(
                buffer,
                Emu((width - pic_w) // 2),
                Emu((height - pic_h) // 2),
                Emu(pic_w),
                Emu(pic_h),
            )
            page.close()
    finally:
        doc.close()
    target = out_dir / output_name(item.name, "pptx")
    deck.save(target)
    return [target]


OP = Op(slug="pdf-to-powerpoint", accepts=frozenset({"pdf"}), run=run, options=NoOptions)
